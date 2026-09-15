# -*- coding: utf-8 -*-
"""
Created on Tue Aug  4 16:30:09 2026

@author: Dilip & Chatgpt
"""

from docx import Document
from datetime import datetime
import os
import random

quote_doc = Document(
    r"C:\Users\HP\OneDrive\Desktop\Unnamed folder\Diary.docx"
)

quotes = []

for para in quote_doc.paragraphs:
    if para.text.strip():
        quotes.append(para.text.strip())   

today = datetime.today()

year = today.strftime("%Y")
month = today.strftime("%B")

base_folder = r"C:\Users\HP\Diary"

folder = os.path.join(base_folder, year, month)

os.makedirs(folder, exist_ok=True)

filename = today.strftime("%Y-%m-%d") + ".docx"

filepath = os.path.join(folder, filename)

mood = None
description = ""


def create_diary():

    doc = Document(
        r"C:\Users\HP\OneDrive\Desktop\Unnamed folder\Sample_teplate.docx"
    )

    doc.add_paragraph(
        "Date : " + today.strftime("%A, %d %B %Y")
    )

    doc.add_heading("Mood", level=2)

    doc.add_paragraph("Mood Rating: Not entered")

    doc.add_paragraph("Feeling: ")

    doc.add_heading("Today's Achievements", level=2)
    doc.add_paragraph("")

    doc.add_heading("What did I learn today?", level=2)
    doc.add_paragraph("")

    doc.add_heading("Three things I am grateful for", level=2)
    doc.add_paragraph("1.\n2.\n3.")

    if quotes:
        quote = random.choice(quotes)

        doc.add_heading("Quote of the Day", level=2)
        doc.add_paragraph(quote)

    doc.add_heading("Tomorrow's Goal", level=2)
    doc.add_paragraph("")

    doc.save(filepath)

    print(f"{filename} created successfully!")
    
    
if os.path.exists(filepath):

    print("Today's diary already exists.")

else:

    create_diary()


























